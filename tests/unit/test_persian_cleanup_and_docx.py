"""Unit tests for Persian text normalization and RTL DOCX export."""

from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

from pdf_translator.adapters.export.docx_exporter import DocxExporter
from pdf_translator.application.persian_cleanup import (
    clean_markdown_persian,
    fix_arabic_chars,
    fix_persian_punctuation,
    fix_punctuation_spacing,
    fix_zwnj_prefixes,
    fix_zwnj_suffixes,
    fix_zwnj_verbs,
    to_english_digits,
    to_persian_digits,
)
from pdf_translator.domain.entities import Page, Project


def test_digit_conversions():
    assert to_persian_digits("Page 123 of 456") == "Page ۱۲۳ of ۴۵۶"
    assert to_english_digits("صفحه ۱۲۳") == "صفحه 123"


def test_arabic_to_persian_chars():
    raw = "يک روز برفی با کيک و قهوة"
    normalized = fix_arabic_chars(raw)
    assert normalized == "یک روز برفی با کیک و قهوه"
    assert "ي" not in normalized
    assert "ك" not in normalized


def test_zwnj_rules():
    # Verbs
    assert fix_zwnj_verbs("می شود و نمی دانم") == "می‌شود و نمی‌دانم"
    
    # Suffixes
    assert fix_zwnj_suffixes("کتاب ها و مقاله های بزرگ تر") == "کتاب‌ها و مقاله‌های بزرگ‌تر"
    
    # Prefixes
    assert fix_zwnj_prefixes("به عنوان مثال یک ابزار بی نظیر است") == "به‌عنوان مثال یک ابزار بی‌نظیر است"


def test_persian_punctuation_and_spacing():
    raw = "چرا ؟ زیرا تست است ;"
    punct = fix_persian_punctuation(raw)
    assert "؟" in punct
    assert "؛" in punct
    assert "?" not in punct
    assert ";" not in punct

    spacing = fix_punctuation_spacing("سلام ، خوبی ؟")
    assert spacing == "سلام، خوبی؟"

    quotes = fix_persian_punctuation('"معماری نرم‌افزار"')
    assert quotes == "«معماری نرم‌افزار»"

    em_dash = fix_persian_punctuation("این مبحث — به ویژه بخش اول — مهم است")
    assert "—" not in em_dash


def test_clean_markdown_persian_preserves_technical_elements():
    md_input = """# فصل 1 : مقدمه ای بر سیستم ها
در این فصل ، ما می دانیم که چگونه سیستم های توزیع شده کار می کنند .
به عنوان مثال "مقیاس پذیری" بسیار با اهمیت است .

1. مورد اول: کتاب ها
2. مورد دوم: بزرگ تر

بررسی کد زیر:
```python
def calculate(x, y):
    # This must remain 100% untouched: x=1, y="test"?
    return x + y
```

فرمول ریاضی: $E = m, c = 2$ و فرمول بلوکی:
$$
\\sum_{i=1}^{n} x_i = y
$$

تگ‌های سیستمی: [[CODE_BLOCK_0]] و [[IMAGE_BLOCK_1]]
لینک: [مستندات اصلی](https://example.com/api?id=12,34)
تصویر: ![نمودار معماری](/api/projects/1/pages/2/images/1)
"""

    cleaned = clean_markdown_persian(md_input)

    # 1. Check Persian normalization applied
    assert "سیستم‌ها" in cleaned
    assert "می‌دانیم" in cleaned
    assert "کار می‌کنند." in cleaned
    assert "به‌عنوان مثال" in cleaned
    assert "«مقیاس پذیری»" in cleaned
    assert "۱. مورد اول: کتاب‌ها" in cleaned
    assert "۲. مورد دوم: بزرگ‌تر" in cleaned

    # 2. Check code block preserved
    assert 'x=1, y="test"?' in cleaned
    assert "def calculate(x, y):" in cleaned

    # 3. Check math preserved
    assert "$E = m, c = 2$" in cleaned
    assert "\\sum_{i=1}^{n} x_i = y" in cleaned

    # 4. Check placeholders and URLs preserved
    assert "[[CODE_BLOCK_0]]" in cleaned
    assert "[[IMAGE_BLOCK_1]]" in cleaned
    assert "https://example.com/api?id=12,34" in cleaned
    assert "/api/projects/1/pages/2/images/1" in cleaned


def test_docx_exporter_rtl_and_vazirmatn(tmp_path: Path):
    exporter = DocxExporter()
    page = Page(
        project_id="test-proj-rtl",
        page_number=3,
        source_text="Test source",
        approved_text="""# مبحث مقیاس‌پذیری
در این بخش موارد زیر بررسی می‌شوند:
* مورد اول: مقیاس عمودی
* مورد دوم: مقیاس افقی

1. گام اول: طراحی معماری
2. گام دوم: بنچ‌مارک

> یک سیستم خوب، پیش‌بینی‌پذیر است.
""",
    )
    out_file = tmp_path / "page_export.docx"
    exporter.export_page(page, out_file)
    assert out_file.exists()
    assert out_file.stat().st_size > 0

    # Inspect XML properties of generated docx
    doc = Document(str(out_file))
    
    # 1. Section bidi
    assert doc.sections[0]._sectPr.find(qn("w:bidi")) is not None

    # 2. Heading uses Vazirmatn and bidi
    h1 = doc.paragraphs[0]
    assert "صفحه ۳" in h1.text
    assert h1._p.get_or_add_pPr().find(qn("w:bidi")) is not None
    assert h1._p.get_or_add_pPr().find(qn("w:jc")).attrib[qn("w:val")] == "start"

    # 3. Check Vazirmatn font set on run
    run_font = h1.runs[0]._r.get_or_add_rPr().find(qn("w:rFonts"))
    assert run_font.attrib[qn("w:cs")] == "Vazirmatn"
    assert h1.runs[0]._r.get_or_add_rPr().find(qn("w:rtl")) is not None


def test_docx_exporter_assemble_book(tmp_path: Path):
    exporter = DocxExporter()
    project = Project(
        id="proj-123",
        title="کتاب سیستم‌های توزیع‌شده",
        source_language="English",
        target_language="Persian",
        total_pages=1,
    )
    page = Page(
        project_id=project.id,
        page_number=1,
        source_text="Source",
        approved_text="ترجمه کامل و روان.",
    )
    out_book = tmp_path / "book_export.docx"
    exporter.assemble_book(project, [page], out_book)
    assert out_book.exists()
    assert out_book.stat().st_size > 0

def test_clean_markdown_persian_normalizes_bold_with_spaces():
    """Verifies that bold delimiters with trailing spaces before closing ** are normalized to valid CommonMark."""
    raw = "**Embedding همیشه مناسب نیست: ** این یک توضیح است."
    cleaned = clean_markdown_persian(raw)
    assert "**Embedding همیشه مناسب نیست:**" in cleaned
    assert not cleaned.startswith("**Embedding همیشه مناسب نیست: **")

def test_clean_markdown_persian_normalizes_consecutive_bold_items_with_hyphens():
    """Verifies that consecutive bold items with trailing spaces and hyphens (**Word **- **Word 2 **-)
    are all cleanly normalized to valid CommonMark bold without corrupting or skipping earlier items.
    """
    raw = "**Star Schema **- **Snowflake Schema **- **Dimensional Modeling **- **One Big Table (OBT) **- فرایند ETL داده‌های سیستم‌های عملیاتی را به schema انتخاب‌شده در Data Warehouse تبدیل می‌کند."
    cleaned = clean_markdown_persian(raw)
    assert "**Star Schema** - " in cleaned
    assert "**Snowflake Schema** - " in cleaned
    assert "**Dimensional Modeling** - " in cleaned
    assert "**One Big Table (OBT)** - " in cleaned
    assert "**Star Schema **" not in cleaned
    assert "**Snowflake Schema **" not in cleaned
    assert "**Dimensional Modeling **" not in cleaned
    assert "**One Big Table (OBT) **" not in cleaned

def test_clean_markdown_persian_user_log_structured_storage_case():
    """Verifies user-reported issue where bold terms with trailing spaces and colons before words
    are correctly normalized with valid CommonMark spacing so they render as bold HTML.
    """
    raw = "**Log-Structured Storage Engines: **داده را در فایل‌های تغییرناپذیر (Immutable Data Files) و معمولاً به‌صورت Append-only می‌نویسند."
    cleaned = clean_markdown_persian(raw)
    assert "**Log-Structured Storage Engines:** داده را" in cleaned
    assert ": **" not in cleaned

def test_clean_markdown_persian_normalizes_bold_without_space_after_colon():
    """Verifies that closing bold tags followed directly by word characters or ZWNJ without spaces are given a space."""
    raw1 = "**Log-Structured Storage Engines:**داده را ذخیره می‌کند"
    cleaned1 = clean_markdown_persian(raw1)
    assert "**Log-Structured Storage Engines:** داده را" in cleaned1

    raw2 = "**Log-Structured Storage Engines:**‌داده را ذخیره می‌کند"
    cleaned2 = clean_markdown_persian(raw2)
    assert "**Log-Structured Storage Engines:** داده را" in cleaned2

def test_clean_markdown_persian_multiple_bold_on_single_line_does_not_corrupt_delimiters():
    """Verifies that multiple bold items on a single line are independently preserved and do not leak spaces into each other."""
    raw = "- **Log-Structured Storage Engines:** داده را در فایل‌های تغییرناپذیر (**Immutable Data Files**) و معمولاً به‌صورت Append-only می‌نویسند."
    cleaned = clean_markdown_persian(raw)
    assert "- **Log-Structured Storage Engines:** داده را در فایل‌های تغییرناپذیر (**Immutable Data Files**) و معمولاً به‌صورت Append-only می‌نویسند." == cleaned

def test_clean_markdown_persian_normalizes_internal_space_before_colon():
    """Verifies that space before colon inside bold delimiters (**Title : **) is cleaned up."""
    raw = "**Log-Structured Storage Engines : **داده"
    cleaned = clean_markdown_persian(raw)
    assert "**Log-Structured Storage Engines:** داده" in cleaned
